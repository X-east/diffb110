"""Produce a one-slide review version from source slides 13 and 14."""
from __future__ import annotations

import json
from pathlib import Path
import sys

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_TICK_LABEL_POSITION, XL_TICK_MARK, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).parent
BLUE='044875';GREY='586570'


def text(slide,value: str,x: float,y: float,w: float,h: float,size: float=18,
         bold: bool=False,color: str='252525',align: PP_ALIGN=PP_ALIGN.LEFT) -> None:
    """Add text using the original deck's visual hierarchy."""
    shape=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    frame=shape.text_frame;frame.word_wrap=True;frame.vertical_anchor=MSO_ANCHOR.MIDDLE
    frame.margin_top=frame.margin_bottom=0;frame.margin_left=frame.margin_right=0
    for j,line in enumerate(value.split('\n')):
        p=frame.paragraphs[0] if j==0 else frame.add_paragraph()
        p.alignment=align;p.space_after=Pt(1)
        run=p.add_run();run.text=line;run.font.name='Arial';run.font.size=Pt(size)
        run.font.bold=bold;run.font.color.rgb=RGBColor.from_string(color)


def rect(slide,x: float,y: float,w: float,h: float,color: str) -> None:
    """Retain the existing blue header and pale conclusion strip."""
    shape=slide.shapes.add_shape(MSO_SHAPE.RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
    shape.fill.solid();shape.fill.fore_color.rgb=RGBColor.from_string(color);shape.line.fill.background()


def main() -> None:
    """Reuse all four original native charts and retain their data exactly."""
    source=ROOT/'中期/谢东宸中期2.pptx'
    p=Presentation(source)
    for idx,slide_id in reversed(list(enumerate(list(p.slides._sldIdLst)))):
        if idx!=12:p.part.drop_rel(slide_id.rId);p.slides._sldIdLst.remove(slide_id)
    slide=p.slides[0]
    charts=[s for s in slide.shapes if s.has_chart]
    assert len(charts)==4
    chart_values=[list(s.chart.series[0].values) for s in charts]
    for shape in list(slide.shapes):
        if not shape.has_chart:
            shape._element.getparent().remove(shape._element)
    rect(slide,0,.27,.65,.18,BLUE);rect(slide,12.13,.27,1.20,.18,BLUE)
    text(slide,'03',.76,.12,.50,.49,28,True)
    text(slide,'课题进展（一）：强化学习生成优化',1.36,.12,10.5,.49,28,False,BLUE)
    text(slide,'强化学习训练与独立生成评价',.70,.98,11.95,.45,24,True)
    text(slide,'观察四项训练指标的变化，选取epoch80冻结模型，并开展500步独立生成评价。',.70,1.58,11.95,.42,17)
    text(slide,'训练过程｜对接、QED、SA与有效率',.80,2.18,6.48,.37,19,True,BLUE)
    text(slide,'epoch80｜独立生成结果',7.66,2.18,5.00,.37,19,True,BLUE)
    text(slide,'→',7.12,2.18,.42,.37,22,True,BLUE,PP_ALIGN.CENTER)
    labels=['对接均值（kcal/mol）','QED均值','SA均值（越低越易合成）','生成有效率（%）']
    positions=[(.73,2.98),(4.05,2.98),(.73,4.76),(4.05,4.76)]
    for i,(shape,(x,y),label) in enumerate(zip(charts,positions,labels)):
        text(slide,label,x+.07,y-.36,3.10,.29,14.5,True,BLUE)
        shape.left=Inches(x);shape.top=Inches(y);shape.width=Inches(3.18);shape.height=Inches(1.35)
        chart=shape.chart;chart.has_title=False;chart.has_legend=False
        grouping=chart.plots[0]._element.find(qn('c:grouping'))
        if grouping is not None and grouping.get('val')=='none':grouping.set('val','standard')
        dataset=CategoryChartData()
        dataset.categories=[category.label for category in chart.plots[0].categories]
        dataset.add_series(chart.series[0].name,chart_values[i])
        selected=[None]*len(chart_values[i]);selected[79]=chart_values[i][79]
        dataset.add_series('epoch80',selected)
        chart.replace_data(dataset)
        marker=chart.series[1].marker;marker.style=XL_MARKER_STYLE.CIRCLE;marker.size=7
        marker.format.fill.solid();marker.format.fill.fore_color.rgb=RGBColor.from_string('B86B43')
        marker.format.line.color.rgb=RGBColor.from_string('B86B43')
        chart.series[1].format.fill.solid()
        chart.series[1].format.fill.fore_color.rgb=RGBColor.from_string('B86B43')
        chart.series[1].format.line.color.rgb=RGBColor.from_string('B86B43')
        for axis in [chart.category_axis,chart.value_axis]:
            axis.tick_labels.font.name='Arial';axis.tick_labels.font.size=Pt(10.5)
            axis.major_tick_mark=XL_TICK_MARK.NONE;axis.minor_tick_mark=XL_TICK_MARK.NONE
        chart.category_axis.tick_label_position=XL_TICK_LABEL_POSITION.LOW
        assert list(chart.series[0].values)==chart_values[i]
    text(slide,'圆点：epoch80检查点',.80,6.17,4.80,.19,11,False,GREY)
    text(slide,'epoch',6.62,6.17,.65,.19,11,False,GREY,PP_ALIGN.RIGHT)
    figure=OUT/'epoch80_docking.png'
    with Image.open(figure) as im:ratio=im.width/im.height
    w=5.02;h=w/ratio
    slide.shapes.add_picture(str(figure),Inches(7.59),Inches(2.66),width=Inches(w),height=Inches(h))
    rows=[['指标','全样本','R1–R5通过'],['分子数','1858','477'],['占2000次尝试','92.90%','23.85%'],['对接均值','−11.29','−8.87']]
    widths=[2.02,1.40,1.57];table=slide.shapes.add_table(4,3,Inches(7.66),Inches(4.88),Inches(sum(widths)),Inches(1.43)).table
    for j,width in enumerate(widths):table.columns[j].width=Inches(width)
    for i,row in enumerate(rows):
        table.rows[i].height=Inches(.38 if i==0 else .35)
        for j,value in enumerate(row):
            cell=table.cell(i,j);cell.text=value;cell.margin_left=Inches(.06);cell.margin_right=Inches(.04)
            cell.margin_top=cell.margin_bottom=0;cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=RGBColor.from_string(BLUE if i==0 else 'F3F6F8' if i%2 else 'FFFFFF')
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment=PP_ALIGN.LEFT if j==0 else PP_ALIGN.CENTER
                for run in paragraph.runs:
                    run.font.name='Arial';run.font.size=Pt(13 if i==0 else 14)
                    run.font.bold=i==0;run.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '252525')
    rect(slide,.70,6.54,11.95,.48,'EDF4F8')
    text(slide,'训练中对接与QED呈改善趋势；独立生成获得1858个分子，其中477个通过结构判据。',.87,6.60,11.61,.34,16.3,True,BLUE)
    text(slide,'独立生成：2000次固定尝试，500步；表中产率均以2000为分母。白点为对接均值，单位kcal/mol。',.70,7.10,11.35,.16,9.8,False,GREY)
    text(slide,'图中两组来自同一epoch80分子池；规则通过仅表示R1–R5全清。分子性质详见下一页。',.70,7.30,11.35,.14,9.8,False,GREY)
    slide.notes_slide.notes_text_frame.text=(
        '本页合并原第13、14页，保留四张原生训练曲线，数据与原稿逐点相等，共89个epoch记录。选取epoch80检查点开展500步独立生成，不把训练曲线各epoch样本看作独立未调优底座对照。'
        '口头讲述：首先介绍训练中对接和QED的趋势，同时展示SA及有效率；随后说明冻结epoch80、固定2000次尝试，保存1858个单片段合法分子；其中477个R1–R5全清。'
        '右侧小提琴的两组都是同一个独立分子池，不能称模型调优前后。对接均值−11.288与−8.87184，筛后均值不等于活性。完整性质放原第15页。'
        '来源：原P13四张chart；analysis/diffgui_epoch80_independent_eval_261010/slide_assets/independent_molecules_joined.csv。未训练、未重新生成或对接、未做PoseBusters。')
    raw=OUT/'merged_raw.pptx';p.save(raw)
    sys.path.insert(0,str(ROOT/'中期/scripts'))
    from set_methods_fonts import set_fonts
    destination=ROOT/'中期/第13-14页_训练与独立生成合并_v1.pptx'
    set_fonts(raw,destination)
    (OUT/'manifest.json').write_text(json.dumps({'source':str(source),'output':str(destination),'native_charts':4,'chart_values_preserved':True,'score_all_mean':-11.2880069968,'score_clean_mean':-8.8718385744,'attempts':2000,'saved':1858,'clean':477},ensure_ascii=False,indent=2))
    print(destination)


if __name__=='__main__':main()
