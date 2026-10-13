"""Append editable PAβN-reference interaction slides to the molecule atlas PPT."""

from __future__ import annotations

import csv
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
ASSET_ROOT = ROOT / "top1_docking_interactions_20260716"
INPUT_PPTX = ROOT / "DiffGui_DiffSBDD_rep1_PBvalid_top50_original_editable.pptx"
OUTPUT_PPTX = ROOT / (
    "DiffGui_DiffSBDD_rep1_PBvalid_top50_original_editable_with_top1_interactions.pptx"
)
LATIN_FONT = "Arial"
EAST_ASIAN_FONT = "Microsoft YaHei"
SLIDE_WIDTH = Inches(13.333333)
SLIDE_HEIGHT = Inches(7.5)


def _set_run_fonts(run: object, size_pt: float, bold: bool = False) -> None:
    """Set Latin and East Asian fonts explicitly on one text run."""
    run.font.name = LATIN_FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(20, 20, 20)
    properties = run._r.get_or_add_rPr()
    for tag, typeface in (
        ("a:latin", LATIN_FONT),
        ("a:ea", EAST_ASIAN_FONT),
        ("a:cs", LATIN_FONT),
    ):
        element = properties.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            properties.append(element)
        element.set("typeface", typeface)


def _add_text(
    slide: object,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    size_pt: float,
    name: str,
    bold: bool = False,
) -> None:
    """Add one centered editable text box."""
    shape = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.add_run()
    run.text = text
    _set_run_fonts(run, size_pt, bold=bold)


def _fit_image(path: Path, left: float, top: float, width: float, height: float) -> tuple:
    """Return a centered aspect-preserving image rectangle in inches."""
    from PIL import Image

    with Image.open(path) as image:
        image_width, image_height = image.size
    scale = min(width / image_width, height / image_height)
    fitted_width = image_width * scale
    fitted_height = image_height * scale
    return (
        Inches(left + (width - fitted_width) / 2),
        Inches(top + (height - fitted_height) / 2),
        Inches(fitted_width),
        Inches(fitted_height),
    )


def _add_picture(slide: object, path: Path, box: tuple, name: str, alt: str) -> None:
    """Add one independent picture object with alt text."""
    picture = slide.shapes.add_picture(str(path), *box)
    picture.name = name
    nodes = picture._element.xpath("./p:nvPicPr/p:cNvPr")
    if nodes:
        nodes[0].set("descr", alt)


def _load_rows() -> dict[str, dict[str, str]]:
    """Load the generated Top1 asset mapping."""
    path = ASSET_ROOT / "top1_interaction_asset_mapping.csv"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    result = {row["model"]: row for row in rows}
    if set(result) != {"DiffGui", "DiffSBDD"}:
        raise ValueError("Expected DiffGui and DiffSBDD Top1 rows")
    return result


def _add_reference_slide(
    presentation: Presentation, model: str, row: dict[str, str]
) -> None:
    """Append one PAβN-reference comparison slide."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)

    title = f"RE-CmeB(8GJL)_PAβN_vs_{model}_mol001"
    _add_text(slide, title, 0.0, 0.0, 4.50, 0.4, 13.0, "interaction_title", bold=False)

    _add_picture(
        slide,
        ASSET_ROOT / Path(row["pymol_png"]).name,
        _fit_image(ASSET_ROOT / Path(row["pymol_png"]).name, 0.0, 0.72, 8.18, 6.05),
        f"{model}_PAβN_mol001_pymol_overlay",
        f"8GJL receptor with docked PAβN and {model} mol001",
    )
    _add_text(
        slide,
        f"8GJL_PAβN_{model}_mol001 | PAβN {float(row['pabn_docking_score']):.3f} | mol001 {float(row['docking_score']):.3f}",
        0.95,
        6.84,
        5.85,
        0.4,
        9.2,
        "docking_score_caption",
    )

    pabn_path = ASSET_ROOT / "PABN_2d.png"
    candidate_path = ASSET_ROOT / f"{model}_mol001_2d.png"
    _add_picture(
        slide,
        pabn_path,
        _fit_image(pabn_path, 7.43, 0.09, 2.37, 1.84),
        f"{model}_PABN_2d",
        "PAβN 2D structure",
    )
    _add_picture(
        slide,
        candidate_path,
        _fit_image(candidate_path, 10.21, 0.40, 1.97, 1.34),
        f"{model}_mol001_2d",
        f"{model} mol001 2D structure",
    )
    _add_text(slide, "PAβN (cyan)", 0.35, 0.42, 1.25, 0.3, 9.5, "pabn_label", bold=True)
    _add_text(
        slide,
        "mol001 (yellow)",
        1.65,
        0.42,
        1.45,
        0.3,
        9.5,
        "candidate_label",
        bold=True,
    )

    interaction_path = ASSET_ROOT / Path(row["interaction_png"]).name
    _add_picture(
        slide,
        interaction_path,
        _fit_image(interaction_path, 7.38, 1.97, 5.35, 4.91),
        f"{model}_mol001_2d_interaction",
        f"Schrodinger 2D interaction diagram for docked {model} mol001 on 8GJL",
    )
    _add_text(
        slide,
        f"8GJL_{model}_mol001_2d_interaction",
        8.06,
        6.84,
        3.99,
        0.4,
        9.2,
        "interaction_caption",
    )


def main() -> None:
    """Append two editable interaction slides and validate their object counts."""
    rows = _load_rows()
    presentation = Presentation(str(INPUT_PPTX))
    if len(presentation.slides) != 4:
        raise ValueError(f"Expected the original 4-slide atlas, found {len(presentation.slides)}")
    for model in ("DiffGui", "DiffSBDD"):
        _add_reference_slide(presentation, model, rows[model])
    if len(presentation.slides) != 6:
        raise ValueError(f"Expected 6 slides after append, found {len(presentation.slides)}")
    for index, slide in enumerate(list(presentation.slides)[4:], start=5):
        pictures = sum(shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes)
        texts = sum(shape.has_text_frame for shape in slide.shapes)
        if pictures != 4 or texts != 5:
            raise ValueError(f"Slide {index}: pictures={pictures}, text_boxes={texts}")
    presentation.save(str(OUTPUT_PPTX))
    print(OUTPUT_PPTX)


if __name__ == "__main__":
    main()
