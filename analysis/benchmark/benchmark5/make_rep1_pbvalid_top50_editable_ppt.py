"""Build a true editable four-slide PPTX from independent molecule assets."""

from __future__ import annotations

import argparse
import csv
import shutil
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt


BASE = Path(__file__).resolve().parent
DEFAULT_ROOT = BASE / "rep1_pbvalid_top50_editable_ppt_20260714"
LATIN_FONT = "Arial"
EAST_ASIAN_FONT = "Microsoft YaHei"
SLIDE_WIDTH = Inches(13.333333)
SLIDE_HEIGHT = Inches(7.5)
GRID_COLUMNS = 5
GRID_ROWS = 5


def _set_run_fonts(run: object, size_pt: float, bold: bool = False) -> None:
    """Set Latin and East Asian typefaces explicitly on one text run."""
    run.font.name = LATIN_FONT
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(20, 20, 20)
    run_properties = run._r.get_or_add_rPr()
    for tag, typeface in (
        ("a:latin", LATIN_FONT),
        ("a:ea", EAST_ASIAN_FONT),
        ("a:cs", LATIN_FONT),
    ):
        element = run_properties.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            run_properties.append(element)
        element.set("typeface", typeface)


def _add_text(
    slide: object,
    text: str,
    left: int,
    top: int,
    width: int,
    height: int,
    size_pt: float,
    bold: bool = False,
    name: str = "text",
) -> None:
    """Add one centered editable text box with explicit script fonts."""
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.name = name
    frame = shape.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    _set_run_fonts(run, size_pt=size_pt, bold=bold)


def _fit_image_box(
    image_path: Path,
    left: int,
    top: int,
    width: int,
    height: int,
) -> tuple[int, int, int, int]:
    """Fit an image in a bounding box without distortion."""
    with Image.open(image_path) as image:
        image_width, image_height = image.size
    scale = min(width / image_width, height / image_height)
    fitted_width = max(1, int(image_width * scale))
    fitted_height = max(1, int(image_height * scale))
    fitted_left = left + (width - fitted_width) // 2
    fitted_top = top + (height - fitted_height) // 2
    return fitted_left, fitted_top, fitted_width, fitted_height


def _set_alt_text(picture: object, description: str) -> None:
    """Set an accessibility description on a picture shape."""
    non_visual = picture._element.xpath("./p:nvPicPr/p:cNvPr")
    if non_visual:
        non_visual[0].set("descr", description)


def _load_mapping(root: Path) -> dict[str, list[dict[str, str]]]:
    """Load and validate original-structure molecule assets."""
    path = root / "original_structure_asset_mapping.csv"
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    grouped: dict[str, list[dict[str, str]]] = {}
    for model in ("DiffGui", "DiffSBDD"):
        selected = [row for row in rows if row["model"] == model]
        selected.sort(key=lambda row: int(row["rank"]))
        if len(selected) != 50:
            raise ValueError(f"{model}: expected 50 assets, found {len(selected)}")
        for row in selected:
            if not Path(row["asset_png"]).exists():
                raise FileNotFoundError(row["asset_png"])
        grouped[model] = selected
    return grouped


def _set_theme_fonts(pptx_path: Path) -> None:
    """Set PowerPoint theme defaults to Arial and Microsoft YaHei."""
    namespace = "http://schemas.openxmlformats.org/drawingml/2006/main"
    with tempfile.NamedTemporaryFile(
        suffix=".pptx", dir=pptx_path.parent, delete=False
    ) as temporary:
        temporary_path = Path(temporary.name)
    try:
        with zipfile.ZipFile(pptx_path, "r") as source_zip, zipfile.ZipFile(
            temporary_path, "w", compression=zipfile.ZIP_DEFLATED
        ) as target_zip:
            for item in source_zip.infolist():
                payload = source_zip.read(item.filename)
                if item.filename == "ppt/theme/theme1.xml":
                    root = etree.fromstring(payload)
                    for font_group in ("majorFont", "minorFont"):
                        group = root.find(f".//{{{namespace}}}{font_group}")
                        if group is None:
                            continue
                        latin = group.find(f"{{{namespace}}}latin")
                        east_asian = group.find(f"{{{namespace}}}ea")
                        complex_script = group.find(f"{{{namespace}}}cs")
                        if latin is not None:
                            latin.set("typeface", LATIN_FONT)
                        if east_asian is not None:
                            east_asian.set("typeface", EAST_ASIAN_FONT)
                        if complex_script is not None:
                            complex_script.set("typeface", LATIN_FONT)
                        for supplemental in group.findall(f"{{{namespace}}}font"):
                            if supplemental.get("script") in {"Hans", "Hant"}:
                                supplemental.set("typeface", EAST_ASIAN_FONT)
                    payload = etree.tostring(
                        root, xml_declaration=True, encoding="UTF-8", standalone=True
                    )
                target_zip.writestr(item, payload)
        shutil.move(temporary_path, pptx_path)
    finally:
        temporary_path.unlink(missing_ok=True)


def _add_molecule_slide(
    presentation: Presentation,
    model: str,
    start_rank: int,
    rows: list[dict[str, str]],
) -> None:
    """Add one slide containing 25 independent molecule images and labels."""
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
    end_rank = start_rank + 24
    _add_text(
        slide,
        f"{model}  Top {start_rank}-{end_rank}",
        Inches(0.15),
        Inches(0.02),
        Inches(13.033333),
        Inches(0.27),
        size_pt=15,
        bold=True,
        name=f"title_{model}_{start_rank}_{end_rank}",
    )

    margin_x = Inches(0.12)
    grid_top = Inches(0.30)
    grid_width = SLIDE_WIDTH - 2 * margin_x
    grid_height = SLIDE_HEIGHT - grid_top - Inches(0.06)
    cell_width = int(grid_width / GRID_COLUMNS)
    cell_height = int(grid_height / GRID_ROWS)
    image_height = int(cell_height * 0.73)
    name_top_offset = int(cell_height * 0.73)
    name_height = int(cell_height * 0.14)
    score_top_offset = int(cell_height * 0.86)
    score_height = int(cell_height * 0.12)

    page_rows = rows[start_rank - 1 : end_rank]
    if len(page_rows) != 25:
        raise ValueError(f"{model} ranks {start_rank}-{end_rank}: expected 25 rows")
    for local_index, row in enumerate(page_rows):
        grid_row, grid_column = divmod(local_index, GRID_COLUMNS)
        cell_left = int(margin_x + grid_column * cell_width)
        cell_top = int(grid_top + grid_row * cell_height)
        image_path = Path(row["asset_png"])
        image_box = _fit_image_box(
            image_path,
            cell_left + Inches(0.05),
            cell_top + Inches(0.01),
            cell_width - Inches(0.10),
            image_height - Inches(0.01),
        )
        picture = slide.shapes.add_picture(str(image_path), *image_box)
        picture.name = f"{model}_mol{int(row['rank']):03d}_structure"
        _set_alt_text(
            picture,
            f"{model} mol{int(row['rank']):03d} original pre-LigPrep structure",
        )
        _add_text(
            slide,
            row["display_name"],
            cell_left,
            cell_top + name_top_offset,
            cell_width,
            name_height,
            size_pt=11.5,
            bold=True,
            name=f"{model}_mol{int(row['rank']):03d}_name",
        )
        _add_text(
            slide,
            f"Docking score = {float(row['docking_score']):.3f}",
            cell_left,
            cell_top + score_top_offset,
            cell_width,
            score_height,
            size_pt=9.2,
            bold=False,
            name=f"{model}_mol{int(row['rank']):03d}_score",
        )


def _validate_presentation(presentation: Presentation) -> None:
    """Verify slide count and independent editable object counts."""
    if len(presentation.slides) != 4:
        raise ValueError(f"Expected 4 slides, found {len(presentation.slides)}")
    for slide_index, slide in enumerate(presentation.slides, 1):
        picture_count = sum(
            shape.shape_type == MSO_SHAPE_TYPE.PICTURE for shape in slide.shapes
        )
        text_count = sum(shape.has_text_frame for shape in slide.shapes)
        if picture_count != 25 or text_count != 51:
            raise ValueError(
                f"Slide {slide_index}: pictures={picture_count}, text_boxes={text_count}"
            )


def main() -> None:
    """Assemble and validate the editable molecule atlas PPTX."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=DEFAULT_ROOT)
    parser.add_argument(
        "--output-name", default="DiffGui_DiffSBDD_rep1_PBvalid_top50_original_editable.pptx"
    )
    args = parser.parse_args()
    root = args.root.resolve()
    root.mkdir(parents=True, exist_ok=True)
    grouped = _load_mapping(root)

    presentation = Presentation()
    presentation.slide_width = SLIDE_WIDTH
    presentation.slide_height = SLIDE_HEIGHT
    for model in ("DiffGui", "DiffSBDD"):
        for start_rank in (1, 26):
            _add_molecule_slide(presentation, model, start_rank, grouped[model])
    _validate_presentation(presentation)
    output_path = root / args.output_name
    presentation.save(output_path)
    _set_theme_fonts(output_path)
    print(output_path)


if __name__ == "__main__":
    main()
