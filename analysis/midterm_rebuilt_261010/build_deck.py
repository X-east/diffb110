"""Rebuild the midterm deck while preserving its first two source slides."""
from __future__ import annotations

import csv
import json
from pathlib import Path
import sys

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION, XL_TICK_LABEL_POSITION, XL_TICK_MARK
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT=Path(__file__).resolve().parents[2]
OUT=Path(__file__).parent
ASSETS=OUT/'assets'
CH2=ROOT/'analysis/midterm_ch2_results_261009'
RL=ROOT/'analysis/diffgui_epoch80_independent_eval_261010/slide_assets'
TRAIN=ROOT/'analysis/midterm_diffgui_slides_261009'
C={'blue':'044875','teal':'2D8B8C','grey':'A7B3BF','ink':'25313B','muted':'5D6B76',
   'pale':'EDF4F8','green':'EAF3EF','orange':'B86B43','line':'DCE4E9'}


def rgb(value: str) -> RGBColor:
    """Convert a palette value for native PowerPoint shapes."""
    return RGBColor.from_string(C.get(value,value))


def read_csv(path: Path) -> list[dict]:
    """Read existing frozen evidence."""
    with path.open(encoding='utf-8-sig') as handle:return list(csv.DictReader(handle))


class Deck:
    """Author native editable text, shapes, tables and charts."""
    def __init__(self) -> None:
        self.pres=Presentation(ROOT/'中期/谢东宸中期2.pptx')
        for slide_id in list(self.pres.slides._sldIdLst)[2:]:
            self.pres.part.drop_rel(slide_id.rId)
            self.pres.slides._sldIdLst.remove(slide_id)
        self.slide=None
        self.inventory=[]

    def text(self,value: str,x: float,y: float,w: float,h: float,size: float=18,
             bold: bool=False,color: str='ink',align: PP_ALIGN=PP_ALIGN.LEFT) -> None:
        """Create a text frame with explicit dimensions and stable typography."""
        shape=self.slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
        frame=shape.text_frame;frame.clear();frame.word_wrap=True
        frame.margin_left=frame.margin_right=Inches(.015)
        frame.margin_top=frame.margin_bottom=0
        frame.vertical_anchor=MSO_ANCHOR.MIDDLE
        for i,line in enumerate(str(value).split('\n')):
            paragraph=frame.paragraphs[0] if i==0 else frame.add_paragraph()
            paragraph.alignment=align;paragraph.space_after=Pt(3)
            run=paragraph.add_run();run.text=line
            run.font.name='Arial';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=rgb(color)

    def card(self,x: float,y: float,w: float,h: float,fill: str='pale') -> None:
        """Draw a quiet card without decorative border stripes."""
        shape=self.slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,Inches(x),Inches(y),Inches(w),Inches(h))
        shape.fill.solid();shape.fill.fore_color.rgb=rgb(fill);shape.line.fill.background()
        shape.adjustments[0]=.08

    def image(self,file: Path,x: float,y: float,w: float,h: float) -> None:
        """Place an existing figure without distortion."""
        with Image.open(file) as im:ratio=im.width/im.height
        iw=min(w,h*ratio);ih=iw/ratio
        self.slide.shapes.add_picture(str(file),Inches(x+(w-iw)/2),Inches(y+(h-ih)/2),width=Inches(iw),height=Inches(ih))

    def page(self,section: str,title: str,intro: str='') -> None:
        """Start a content page and register its one central question."""
        self.slide=self.pres.slides.add_slide(self.pres.slide_layouts[6])
        self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=rgb('FFFFFF')
        self.text(section,.60,.28,11.8,.30,13,True,'blue')
        self.text(title,.60,.84,12.05,.72,29,True)
        if intro:self.text(intro,.60,1.62,12.05,.45,16.5,False,'muted')
        self.text(str(len(self.pres.slides)),12.20,7.12,.52,.22,11,False,'muted',PP_ALIGN.RIGHT)
        self.inventory.append({'page':len(self.pres.slides),'section':section,'title':title})

    def section(self,number: str,title: str,subtitle: str) -> None:
        """Create a chapter divider with a short conceptual transition."""
        self.slide=self.pres.slides.add_slide(self.pres.slide_layouts[6])
        self.slide.background.fill.solid();self.slide.background.fill.fore_color.rgb=rgb('blue')
        self.text(number,.80,1.12,3.1,1.7,90,True,'FFFFFF')
        self.text(title,.88,3.21,11.6,.9,42,True,'FFFFFF')
        self.text(subtitle,.92,4.43,11.2,.75,22,False,'D7E6EF')
        self.text(str(len(self.pres.slides)),12.15,7.11,.55,.25,11,False,'FFFFFF',PP_ALIGN.RIGHT)
        self.inventory.append({'page':len(self.pres.slides),'title':title})

    def takeaway(self,value: str) -> None:
        """State the present finding instead of repeating future work."""
        self.card(.60,6.42,12.05,.52)
        self.text(value,.80,6.49,11.65,.35,17,True,'blue')

    def foot(self,value: str) -> None:
        """Keep only material denominators and comparison limits on canvas."""
        self.text(value,.63,7.02,11.35,.37,10.5,False,'muted')

    def notes(self,value: str) -> None:
        """Retain evidence sources and limitations in speaker notes."""
        self.slide.notes_slide.notes_text_frame.text=value

    def table(self,headers: list[str],rows: list[list],x: float,y: float,widths: list[float],
              row_h: float=.47,font: float=16,highlight: int=-1) -> None:
        """Add a native editable table with minimal emphasis."""
        table=self.slide.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),
                                         Inches(sum(widths)),Inches(row_h*(len(rows)+1))).table
        for i,width in enumerate(widths):table.columns[i].width=Inches(width)
        for j,row in enumerate([headers]+rows):
            table.rows[j].height=Inches(row_h)
            for i,value in enumerate(row):
                cell=table.cell(j,i);cell.text=str(value)
                cell.margin_left=Inches(.08);cell.margin_right=Inches(.05)
                cell.margin_top=cell.margin_bottom=Inches(.025);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
                fill='blue' if j==0 else 'DDEBF4' if j-1==highlight else 'F2F5F7' if j%2 else 'FFFFFF'
                cell.fill.solid();cell.fill.fore_color.rgb=rgb(fill)
                for paragraph in cell.text_frame.paragraphs:
                    paragraph.alignment=PP_ALIGN.LEFT if i==0 else PP_ALIGN.CENTER
                    for run in paragraph.runs:
                        run.font.name='Arial';run.font.size=Pt(font if j else font-1)
                        run.font.bold=j==0 or j-1==highlight
                        run.font.color.rgb=rgb('FFFFFF' if j==0 else 'blue' if j-1==highlight else 'ink')

    def chart(self,categories: list[str],series: list[tuple[str,list[float],str]],x: float,y: float,w: float,h: float,
              kind: XL_CHART_TYPE=XL_CHART_TYPE.LINE,ymin: float|None=None,ymax: float|None=None,
              values: bool=False,legend: bool=False,fmt: str='0.0') -> None:
        """Use editable charts for native PowerPoint chart types."""
        data=CategoryChartData();data.categories=categories
        for name,vals,_ in series:data.add_series(name,vals)
        chart=self.slide.shapes.add_chart(kind,Inches(x),Inches(y),Inches(w),Inches(h),data).chart
        chart.has_title=False
        chart.has_legend=legend
        if legend:
            chart.legend.position=XL_LEGEND_POSITION.TOP;chart.legend.include_in_layout=False
            chart.legend.font.name='Arial';chart.legend.font.size=Pt(13)
        chart.chart_style=10
        for target in [chart.category_axis,chart.value_axis]:
            target.major_tick_mark=XL_TICK_MARK.NONE;target.minor_tick_mark=XL_TICK_MARK.NONE
            target.tick_labels.font.name='Arial';target.tick_labels.font.size=Pt(12)
            target.tick_labels.font.color.rgb=rgb('muted')
        chart.category_axis.tick_label_position=XL_TICK_LABEL_POSITION.LOW
        chart.value_axis.tick_labels.number_format=fmt
        if '' in categories:
            axis=chart.category_axis._element
            skip=OxmlElement('c:tickLblSkip');skip.set('val','1');axis.append(skip)
        if ymin is not None:chart.value_axis.minimum_scale=ymin
        if ymax is not None:chart.value_axis.maximum_scale=ymax
        chart.value_axis.has_major_gridlines=True
        chart.value_axis.major_gridlines.format.line.color.rgb=rgb('line')
        chart.value_axis.major_gridlines.format.line.width=Pt(.5)
        for idx,(_,_,color) in enumerate(series):
            obj=chart.series[idx]
            obj.format.fill.solid();obj.format.fill.fore_color.rgb=rgb(color)
            obj.format.line.color.rgb=rgb(color);obj.format.line.width=Pt(1.8)
        if values:
            for plot in chart.plots:
                plot.has_data_labels=True;plot.data_labels.position=XL_LABEL_POSITION.OUTSIDE_END
                plot.data_labels.font.name='Arial';plot.data_labels.font.size=Pt(13)
                plot.data_labels.number_format=fmt


def main() -> None:
    """Reconstruct the approved narrative from existing verified evidence."""
    d=Deck()
    ch2=json.loads((CH2/'slide_data.json').read_text())
    train=json.loads((TRAIN/'slide_data.json').read_text())['training']
    trained=json.loads((RL/'slide_data.json').read_text())
    comparison=json.loads((ROOT/'analysis/diffgui_candidate_yield_comparison_261010/summary.json').read_text())
    display=json.loads((ASSETS/'diffgui_display_stats.json').read_text())
    f=lambda v,n=2:f'{float(v):.{n}f}'.replace('-','−')

    d.section('01','研究背景与目标','从RE-CmeB候选分子需求，走向口袋条件下的生成模型调优')
    d.page('01 研究背景','RE-CmeB：候选抑制剂设计的研究需求','围绕细菌外排泵探索候选分子，为后续筛选与验证提供资源。')
    d.card(.65,2.30,5.20,3.55)
    d.text('空肠弯曲杆菌',.90,2.56,4.6,.42,25,True,'blue')
    d.text('食源性病原体\n可引起胃肠炎与肠外感染\n相关抗生素耐药问题受到关注',.90,3.18,4.6,1.45,22)
    d.text('研究关注：RE-CmeB外排泵',.90,5.03,4.6,.38,20,True)
    d.image(ASSETS/'protein.png',6.15,2.30,3.15,3.08)
    d.image(ASSETS/'pabn.png',9.53,2.62,3.0,2.35)
    d.text('RE-CmeB结构',6.25,5.45,2.95,.30,15,False,'muted',PP_ALIGN.CENTER)
    d.text('已有抑制剂示例：PAβN',9.42,5.45,3.20,.30,15,False,'muted',PP_ALIGN.CENTER)
    d.takeaway('已有外排泵抑制剂仍存在应用限制，需要继续探索RE-CmeB候选分子。')
    d.foot('背景与图像沿用原稿；S. Saad Shamsi等，IntechOpen，2023。')
    d.notes('本页是靶点需求，不从疾病背景直接推断现有生成模型不适配。生物学背景与结构/PAβN图均来自用户原稿P4，未引入新的活性结论。')

    d.page('01 研究背景','口袋条件三维扩散模型：生成基础与调优空间','利用口袋几何信息生成原子类型与坐标，再通过任务反馈考察候选质量。')
    d.image(ASSETS/'diffusion.png',.70,2.12,11.90,2.15)
    d.card(.70,4.56,5.75,1.42);d.card(6.85,4.56,5.75,1.42)
    d.text('生成基础',.94,4.75,5.20,.35,22,True,'blue')
    d.text('口袋条件化生成\n直接建模分子三维结构',.94,5.20,5.20,.60,19)
    d.text('调优空间',7.09,4.75,5.20,.35,22,True,'blue')
    d.text('目标性质需要定向反馈\n结构与完整性需要综合评价',7.09,5.20,5.20,.60,19)
    d.takeaway('以通用生成能力为基础，进一步探索面向RE-CmeB需求的模型调优。')
    d.foot('图为扩散生成机制示意；逐步去噪不等同于目标性质已经得到优化。')

    d.page('01 研究背景','强化学习与偏好优化：两类调优信号','在同一靶点条件下，分别从性质反馈与结构优劣监督探索生成优化。')
    for x,title,subtitle,items,color in [(.70,'强化学习','在线性质反馈',['生成分子','计算性质反馈','更新生成策略'],'blue'),
                                       (6.85,'偏好优化','离线优劣监督',['构建优劣分子对','学习相对偏好','更新生成模型'],'teal')]:
        d.card(x,2.30,5.75,3.60)
        d.text(title,x+.23,2.55,5.20,.47,28,True,color)
        d.text(subtitle,x+.23,3.10,5.20,.34,19,True)
        for j,item in enumerate(items):
            d.text(f'{j+1:02d}',x+.24,3.64+j*.62,.63,.39,21,True,color)
            d.text(item,x+1.05,3.64+j*.62,4.10,.39,22)
    d.takeaway('两种策略并列探索，分别引入性质反馈和结构偏好。')
    d.foot('方法背景沿用原稿：Loeffler等，J Cheminform，2024；Gu等，AliDiff，2024。')
    d.notes('本页概念背景为两条方法线提供动机，不宣称强化学习与偏好优化组成已验证串行管线。引用沿用原稿；具体项目目标和实现见后续方法页。')

    d.page('01 研究目标','共同目标：提高符合既定要求的候选产出','在RE-CmeB口袋条件下，考察模型调优对生成质量与候选资源的影响。')
    d.card(.70,2.22,11.90,.96)
    d.text('核心问题：相同生成尝试预算，能否获得更多满足综合条件的候选？',.96,2.43,11.40,.54,24,True,'blue')
    d.card(.70,3.57,5.75,1.49);d.card(6.85,3.57,5.75,1.49,'green')
    d.text('课题一｜DiffGui强化学习',.96,3.81,5.16,.37,23,True,'blue')
    d.text('引入对接、QED、SA与多样性反馈',.96,4.38,5.16,.39,19)
    d.text('课题二｜DiffSBDD结构偏好',7.11,3.81,5.16,.37,23,True,'teal')
    d.text('引入匹配分子对的结构优劣监督',7.11,4.38,5.16,.39,19)
    d.text('共同评价：生成有效性  ·  结构与完整性  ·  分子性质  ·  对接与候选产率',.86,5.48,11.55,.55,20,True)
    d.takeaway('为RE-CmeB候选分子的后续优选、结合模式分析与验证提供支持。')
    d.foot('两条路线分别与各自未调优底座比较；计算筛选条件不等同于实验活性判据。')

    d.section('02','研究思路与方法','分别构建调优信号，通过冻结检查点的生成评价形成证据')
    d.page('02 方法（一）｜强化学习','性质评价转为反馈，持续更新DiffGui生成策略','以RE-CmeB口袋为条件，在生成—评价—更新闭环中开展多指标强化学习。')
    cards=[('口袋条件生成','DiffGui生成原子类型与坐标'),('多维性质评价','对接、QED、SA与指纹多样性'),
           ('批次反馈构建','指标排序与正态分数变换'),('PPO策略更新','重新采样，持续评价')]
    for j,(title,body) in enumerate(cards):
        x=.70+j*3.06;d.card(x,2.37,2.75,2.96)
        d.text(f'{j+1:02d}',x+.19,2.59,.71,.49,29,True,'blue')
        d.text(title,x+.19,3.29,2.37,.62,21,True)
        d.text(body,x+.19,4.13,2.37,.82,19)
        if j<3:d.text('→',x+2.78,3.63,.29,.4,22,True,'blue')
    d.text('验收：冻结检查点 → 独立生成 → 与未调优DiffGui采用同一评价和筛选口径',.86,5.68,11.55,.51,20,True)
    d.takeaway('训练反馈用于更新策略，独立生成与候选产率用于检查实际产出。')
    d.foot('方法来源：现有DiffGui PPO流程；训练评分与重新搜索对接分别记录。')
    d.notes('批内排序、正态分数变换和PPO流程沿用原稿与现有DiffGui训练实现。独立生成指冻结权重后重新采样，不指完全独立于训练目标的生物学验证。')

    d.page('02 方法（二）｜结构偏好','结构匹配的优劣分子对，形成端点级偏好监督','以DiffSBDD为基础，在固定口袋条件下学习相对结构偏好。')
    cards=[('结构评价','以经验规则区分优劣端点'),('匹配偏好对','匹配节点数与元素组成'),
           ('ESRPO训练','当前与参考模型的相对评分'),('检查点生成','结构、完整性、性质与对接评价')]
    for j,(title,body) in enumerate(cards):
        x=.70+j*3.06;d.card(x,2.37,2.75,2.96,'green' if j==3 else 'pale')
        d.text(f'{j+1:02d}',x+.19,2.59,.71,.49,29,True,'teal')
        d.text(title,x+.19,3.29,2.37,.62,21,True)
        d.text(body,x+.19,4.13,2.37,.82,19)
        if j<3:d.text('→',x+2.78,3.63,.29,.4,22,True,'teal')
    d.text('评价规则涵盖大环、复杂环系、小环连接与碳比例；同时关注原始单片段产出。',.86,5.60,11.55,.60,18.5)
    d.takeaway('偏好拟合与生成质量分别评价，并明确未调优底座和各检查点的身份。')
    d.foot('ESRPO为端点去噪代理偏好目标；x-only是损失通道设置，生成仍包含原子类型。')
    d.notes('训练结构标签R1–R4，验收R1–R5；R5是组成护栏，不能从结构匹配训练直接推断其被优化。x-only为坐标损失权重1、类型损失权重0，不是固定生成原子类型。')

    d.section('03','课题进展','分别回答：调优前后生成改变了什么，留下了哪些候选')
    d.page('03 进展（一）｜DiffGui强化学习','五模型横评：选择强化学习底座','比较RE-CmeB口袋条件下的生成表现，并结合已有训练接口确定探索底座。')
    d.image(TRAIN/'benchmark_docking_violin.png',.62,2.24,6.13,3.41)
    d.table(['模型','合法产率%','对接均值','QED','SA'],[
        ['DiffGui','86.80','−10.79','0.623','3.94'],['DiffSBDD','66.90','−9.49','0.607','4.28'],
        ['GCDM-SBDD','51.90','−6.63','0.439','4.33'],['PMDM','49.63','−8.72','0.528','5.73'],
        ['TargetDiff','45.37','−7.85','0.326','5.50']],6.97,2.49,[1.50,1.14,1.16,.86,.87],.51,15.2,0)
    d.text('DiffGui在本批横评中的合法产率、对接与QED表现较好。',.90,5.86,11.3,.31,19)
    d.takeaway('结合生成表现与训练接口，选用DiffGui开展多指标强化学习探索。')
    d.foot('500步；每模型3次重复、每次1000尝试。对接图表为历史成功记录的非正子集，单位kcal/mol。')
    d.notes('来源analysis/midterm_diffgui_slides_261009和同批500步rep4/5/6横评。对接均值与QED/SA聚合方法在源表记录。本页底座横评数−10.79不是下一页2000对照重新补氢对接的−10.81，也不是200步早期训练−9.89。普通正分和可疑分在源表保留。')

    d.page('03 进展（一）｜DiffGui强化学习','强化学习训练与独立生成的前后比较','选取epoch80冻结模型，500步生成；两组各2000次尝试，采用同一新对接协议。')
    d.text('训练观察｜截至epoch80',.75,2.18,4.21,.30,19,True,'blue')
    tr=[r for r in train if int(r['epoch'])<=80]
    labels=[str(r['epoch']) if int(r['epoch']) in (1,20,40,60,80) else '' for r in tr]
    d.text('训练对接均值（kcal/mol）',.80,2.70,4.21,.27,15,True)
    d.chart(labels,[('训练对接',[r['docking_raw'] for r in tr],'blue')],.72,3.04,4.20,1.10,ymin=-12,ymax=-8)
    d.text('训练QED均值',.80,4.29,4.21,.27,15,True)
    d.chart(labels,[('QED',[r['qed_raw'] for r in tr],'teal')],.72,4.61,4.20,1.08,ymin=.5,ymax=.75)
    d.text('epoch',4.18,5.79,.65,.23,12,False,'muted',PP_ALIGN.RIGHT)
    d.text('独立生成｜未调优底座 vs epoch80',5.36,2.18,7.20,.30,19,True,'blue')
    d.image(ASSETS/'diffgui_matched_before_after.png',5.34,2.64,7.25,2.08)
    d.table(['独立生成指标','未调优','epoch80'],[
        ['保存合法产率','86.70%','92.90%'],['单片段R1–R5通过产率','16.00%','23.85%'],
        ['所有保存分子对接均值','−10.81','−11.29'],
        ['规则通过子集QED均值',f(display[0]['clean_qed_mean'],3),f(display[1]['clean_qed_mean'],3)],
        ['规则通过子集SA均值',f(display[0]['clean_sa_mean']),f(display[1]['clean_sa_mean'])]],
        5.48,4.73,[3.53,1.65,1.65],.27,13.5)
    d.takeaway('独立生成的结构规则通过产出增加；进一步用综合筛选检查候选价值。')
    d.foot('新对接LF＋加氢、同一UniDock/vina/fast协议；生成种子不同，为历史批次比较。训练分数不与新对接直接混比。')
    d.notes('原P13/P14合并。本页训练曲线截到80；独立生成2000attempts原模型为历史最早两次重复合并1734条，epoch80为1858条。完整protocol相等，ID65fd181dd019ee42。两个小提琴分组分别为所有保存分子及R1–R5全清子集，颜色代表模型身份。白点均值，清洁子集新对接均值−8.765与−8.872。比较未做同种子配对重复，因此仅描述当前阶段产出。')

    d.page('03 进展（一）｜DiffGui强化学习','相同生成预算下，综合候选产出增加','沿用同一结构与性质筛选条件，再检查不同对接阈值下的候选产率。')
    d.text('逐级筛选｜每组2000次尝试',.76,2.23,5.50,.32,19,True,'blue')
    d.table(['阶段','未调优','epoch80'],[
        ['保存且可解析',1734,1858],['单片段＋R1–R5通过',320,477],['再通过性质条件',248,436],
        ['再要求对接≤−10',70,118],['上行去重结构',69,117],['上行骨架数',63,100]],.75,2.81,[3.13,1.26,1.26],.43,15.5,3)
    d.text('对接阈值敏感性｜候选产率（%）',6.88,2.23,5.65,.32,19,True,'blue')
    values=comparison['candidate_thresholds']
    series=[]
    for name,color,label in [('Base first two repeats','grey','未调优'),('Epoch80','blue','epoch80')]:
        series.append((label,[r['yield_pct'] for r in values if r['group']==name],color))
    d.chart(['≤−8','≤−9','≤−10','≤−11','≤−12'],series,6.82,2.79,5.72,2.78,ymin=0,ymax=18,legend=True)
    d.text('性质通过子集中，对接≤−10比例：28.2% → 27.1%。\n当前增益主要体现为通过综合条件的候选数量增加。',6.94,5.68,5.55,.58,15.5)
    d.takeaway('以−10示例阈值计，候选产率3.5%→5.9%；去重候选69→117个。')
    d.foot('性质条件：QED≥0.5、标准SA≤5.5、HA10–30、至少1环。阈值网格是计算敏感性分析，不是活性判据。')
    d.notes('来源analysis/diffgui_candidate_yield_comparison_261010。五档候选数原模型183/139/70/18/2，epoch80为310/205/118/29/7。性质通过子集对接均值−9.039→−8.955，没有改善；不得用总体−10.813→−11.288替代筛后结论。训练成本与生成GPU时间未匹配，等尝试不等于等总计算成本。R1–R5是经验筛选规则；SA越低通常更易合成，仅为代理，不是合成保证。')

    d.page('03 进展（一）｜DiffGui强化学习','已有候选集合：展示不同骨架的代表实例','在−10示例阈值下获得117个去重候选、100个骨架；以下为进一步复核的实例。')
    candidates=trained['candidates'] if 'candidates' in trained else trained.get('representatives')
    if candidates is None:candidates=read_csv(RL/'representative_candidates.csv')
    for j,r in enumerate(candidates[:4]):
        x=.73+(j%2)*6.14;y=2.26+(j//2)*1.93
        d.card(x,y,5.82,1.72,'F6F8FA')
        file=RL/(r.get('asset',f'candidate_{j+1}')+'.png')
        if not file.exists():file=RL/f'candidate_{j+1}.png'
        d.image(file,x+.13,y+.21,2.60,1.25)
        d.text(f'候选 {j+1}｜{r.get("record_id","")}',x+2.89,y+.23,2.65,.28,17,True,'blue')
        d.text(f'对接 {f(r["score"])} kcal/mol\nQED {f(r["qed"],3)}  ·  SA {f(r["sa"])}\nLogP {f(r["logp"])}',x+2.89,y+.72,2.65,.69,15.5)
    d.takeaway('第一部分已形成候选分子池，可继续开展结构复核、结合模式分析与验证。')
    d.foot('实例来自epoch80独立生成；满足既有综合条件，按对接排序并保留不同Murcko骨架。计算评分不等同活性。')
    d.notes('实际source record与源SDF见analysis/diffgui_epoch80_independent_eval_261010/slide_assets/representative_candidates.csv和sdf。此处是四个展示实例，不是仅生成四个候选，也不是所有候选都具备可合成性或实验活性。没有运行PoseBusters。')

    d.page('03 进展（二）｜DiffSBDD结构偏好','从结构评价视角，选择偏好优化底座','结合生成结构观察与ChEMBL参照统计，补充结构筛选和底座选择。')
    d.card(.71,2.13,11.90,.81)
    d.text('经验判据：R1 ≥7元环；R2 连通环系>3环；R3 四元环；R5 碳比例>85%。\nR4 三元环含杂原子、多位点外接，或与其他环形成螺/稠合连接。',.89,2.23,11.53,.60,15.5)
    d.image(CH2/'benchmark_docking_violin.png',.66,3.05,6.52,2.38)
    d.table(['模型','LF通过产率%','筛前对接','规则通过对接'],[
        [r['model'],f(r['lf_clean_yield']),f(r['dock_mean']),f(r['dock_clean_mean'])] for r in ch2['benchmark']],7.37,3.16,[1.40,1.36,1.21,1.27],.39,13.7,1)
    d.image(CH2/'benchmark_compliant.png',.98,5.40,1.90,.64)
    d.text('规则通过实例',2.99,5.57,2.65,.30,16,True,'teal')
    d.image(CH2/'benchmark_rule_hit.png',7.08,5.40,1.90,.64)
    d.text('R4命中实例（着色）',9.09,5.57,3.21,.30,16,True,'orange')
    d.takeaway('结合结构表现、目标性质与训练条件，选用DiffSBDD开展结构偏好探索。')
    d.foot('同批500步横评，各3000尝试；LF为最大片段，不能直接称完整候选。经验规则不等于化学不可能性判断。')
    d.notes('并列路线的结构评价视角，不否定第一部分DiffGui选型。规则依据是生成观察与ChEMBL36参照经验统计，不是全部阈值均为文献共识。图表只用历史成功非正对接。颜色此页代表同池筛前/规则通过，图例已明示；不作方法前后比较。')

    d.page('03 进展（二）｜DiffSBDD结构偏好','结构匹配偏好训练：从真实优劣端点出发','匹配节点数与元素组成，采用x-only端点偏好目标开展训练。')
    d.card(.75,2.32,7.07,3.50)
    d.text('真实匹配对｜均为14个重原子',.97,2.55,6.62,.37,22,True,'blue')
    d.image(ASSETS/'pair_winner.png',.99,3.18,3.03,1.45)
    d.image(ASSETS/'pair_loser.png',4.49,3.18,3.03,1.45)
    d.text('优选：R1–R4全清',.99,4.87,3.03,.34,17,True,'teal',PP_ALIGN.CENTER)
    d.text('劣选：命中R3',4.49,4.87,3.03,.34,17,True,'orange',PP_ALIGN.CENTER)
    d.text('节点数与元素组成匹配，降低组成差异对偏好的干扰。',.99,5.42,6.56,.24,15)
    d.text('前10,000次更新｜训练记录',8.15,2.35,4.44,.36,19,True,'blue')
    loss=ch2['losses']
    d.chart([f'{r["updates"]:,}' for r in loss],[('训练loss',[r['loss'] for r in loss],'blue')],8.07,2.96,4.56,1.94,ymin=.68,ymax=.70,fmt='0.000')
    d.text('学习率 1e−5\n损失通道 x-only\n每2000次更新保存的5个记录点',8.21,5.06,4.19,.75,17)
    d.takeaway('已完成偏好训练；通过检查点生成分别考察结构产出与综合表现。')
    d.foot('训练图不是逐步连续loss；不以loss下降代替生成效果。x-only不意味着固定生成原子类型。')
    d.notes('前10000次更新是展示范围，不代表完整epoch仅10000updates。真实配对winner r02_s0786_m0009，loser r02_s2431_m0003，来自原稿方法配图。训练R1–R4，验收R1–R5。按用户要求只写训练，不写train/val/test对数或准确率。')

    d.page('03 进展（二）｜DiffSBDD结构偏好','未调优底座与检查点：结构产出对比','加入真实未调优DiffSBDD历史参照，区分完整单片段与最大片段评价。')
    d.chart(['最大片段规则通过','原始单片段规则通过'],[
        ('未调优历史参照',[30.1,27.7],'grey'),('step4000',[55.6,41.9],'teal'),
        ('step10000',[62.6,41.3],'blue')],.74,2.39,7.15,3.44,
        kind=XL_CHART_TYPE.COLUMN_CLUSTERED,ymin=0,ymax=75,values=True,legend=True)
    d.card(8.26,2.48,4.34,3.28)
    d.text('历史结构参照',8.49,2.73,3.83,.37,22,True,'blue')
    d.text('未调优：3次横评，共3000尝试\n检查点：各1000尝试\n均以全部尝试为产率分母',8.49,3.34,3.83,1.00,17)
    d.text('单片段规则通过产率\n27.7% → 41.9%（step4000）',8.49,4.79,3.83,.65,20,True)
    d.takeaway('结构规则通过产出发生变化；进一步检查完整性、性质和对接的兼顾。')
    d.foot('生成种子及后处理未完全配对，属于历史参照；LF增加不能全部解释为完整候选增加。')
    d.notes('真实未调优基线来自五模型rep4/5/6，原始单片段R1–R5全清831/3000、LF全清903/3000。step4000为419/1000与556/1000；step10000为413/1000与626/1000。禁止采用260901被错标Original的已训练x-only305/1000批次。严格同生成配置、同种子与同对接的前后候选比较尚未完成。本页是历史描述，不作因果或显著性结论。')

    d.page('03 进展（二）｜DiffSBDD结构偏好','结构变化与综合表现，需要同时兼顾','比较step4000与step10000，检查性质、完整性以及规则通过子集的对接表现。')
    d.table(['同一评价范围','step4000','step10000'],[
        ['原始单片段规则通过产率','41.9%','41.3%'],['LF规则通过产率','55.6%','62.6%'],
        ['原始多片段／全部尝试','23.2%','31.2%'],['规则通过子集QED中位','0.561','0.458'],
        ['规则通过子集对接中位','−8.461','−8.012']],.78,2.46,[4.33,1.53,1.53],.58,17)
    d.card(8.54,2.47,4.06,3.45,'green')
    d.text('阶段发现',8.77,2.76,3.59,.38,23,True,'teal')
    d.text('LF通过产率增加\n完整单片段产率变化较小\n性质与完整性仍需兼顾',8.77,3.40,3.59,1.48,20)
    d.text('step4000作为实例展示检查点，\n不将其称为已验证最优。',8.77,5.12,3.59,.53,15.5,False,'muted')
    d.takeaway('结构偏好带来阶段性结构变化，综合候选表现仍有进一步优化空间。')
    d.foot('每检查点1000尝试；QED为LF规则全清子集，对接为历史成功非正子集；子集随检查点变化。')
    d.notes('Source checkpoint_metrics.csv。尚无真实未调优底座与各检查点同准备流程的新对接，因此不拼接不同准备protocol声称亲和力增益。当前已有结构前后历史参照，综合候选产率需补同口径比较。计算分数不是活性验证。表里的碎片/尝试包含原始可读SDF记录，无效尝试不进入性质中位数。')

    d.page('03 进展（二）｜DiffSBDD结构偏好','已有结构实例：继续复核与候选优选','从step4000原始单片段分子中，筛选不同骨架的规则通过实例。')
    for j,r in enumerate(ch2['candidates'][:4]):
        x=.73+(j%2)*6.14;y=2.26+(j//2)*1.93
        d.card(x,y,5.82,1.72,'F6F8FA');d.image(CH2/(r['asset']+'.png'),x+.13,y+.21,2.60,1.25)
        d.text(f'实例 {j+1}｜SDF {r["raw_indices"][0]}',x+2.89,y+.23,2.65,.28,17,True,'teal')
        d.text(f'对接 {f(r["score"])} kcal/mol\nQED {f(r["qed"],3)}\n重原子数 {r["ha"]}',x+2.89,y+.72,2.65,.69,16)
    d.takeaway('已有可供进一步复核的实例；后续完善模型综合表现并推进候选优选。')
    d.foot('筛选：原始单片段、R1–R5全清、HA10–30、含环、QED≥0.5；对接为历史评分，不直接与第一部分排名。')
    d.notes('四例来自真实raw step4000SDF，序号815/192/845/451，0-based；规范SMILES与relaxed docking输入核对。全部来源CSV见candidate_examples.csv。这是实例展示，不是相对Original的候选产率结果，未补算SA或新对接，不声称活性。')

    d.page('03 研究总结','阶段贡献与主要难点','围绕靶点口袋条件生成，形成两类调优探索和可追溯的候选评价。')
    d.card(.74,2.33,5.75,3.59)
    d.text('阶段贡献',.99,2.59,5.23,.40,27,True,'blue')
    d.text('策略探索\n分别引入强化学习与结构偏好优化\n\n候选资源\n形成可供后续优选与验证的分子集合',.99,3.24,5.23,2.18,20)
    d.card(6.86,2.33,5.75,3.59,'green')
    d.text('主要难点',7.11,2.59,5.23,.40,27,True,'teal')
    d.text('反馈与判据\n合理构建奖励及结构偏好监督\n\n综合评价\n兼顾结构、性质、完整性与候选产出',7.11,3.24,5.23,2.18,20)
    d.takeaway('第一部分已有候选池；第二部分继续完善综合表现，两者均推进候选验证。')
    d.foot('当前为计算阶段结果；候选新颖性、可合成性与活性需要进一步核查和验证。')

    d.section('04','后续安排','以当前结果为依据，分别推进可重复评价、模型完善与候选验证')
    d.page('04 后续安排','分别推进两部分研究，形成候选筛选与验证闭环','延续已有结果，优先补齐支撑结论的前后对照与候选证据。')
    for x,title,stage,steps,color in [(.73,'一｜DiffGui强化学习','已有独立生成候选池',['同种子配对重复，确认候选产率与成本','优选不同骨架，分析结合模式','推进可合成性核查与后续验证'],'blue'),
        (6.86,'二｜DiffSBDD结构偏好','结构表现已改变，综合性质仍需完善',['完善偏好对与训练策略','补齐同口径底座／检查点候选比较','生成优选候选，分析结合模式并验证'],'teal')]:
        d.card(x,2.32,5.75,3.75,'pale' if color=='blue' else 'green')
        d.text(title,x+.24,2.59,5.22,.43,24,True,color)
        d.text(stage,x+.24,3.20,5.22,.54,18,True)
        for j,step in enumerate(steps):
            d.text(f'{j+1}',x+.25,4.00+j*.56,.35,.36,19,True,color)
            d.text(step,x+.77,4.00+j*.56,4.57,.44,17.5)
    d.takeaway('将模型调优的阶段结果，转化为可复核的候选优选与后续验证依据。')
    d.notes('未来模型和奖励机制调整需要既有代码或文献依据及验收消融，不能仅凭此页泛化方向直接上线。第一部分已有候选产率历史比较但未计训练摊销成本；第二部分需匹配真实官方底座并核对生成后处理与新对接。此处为安排，不表示已完成实验验证。')

    d.slide=d.pres.slides.add_slide(d.pres.slide_layouts[6])
    d.slide.background.fill.solid();d.slide.background.fill.fore_color.rgb=rgb('blue')
    d.text('请各位老师批评指正',.95,2.50,11.43,1.00,44,True,'FFFFFF',PP_ALIGN.CENTER)
    d.text('谢东宸  ·  药学\n指导老师：杨龙华\n2026.10.12',3.15,4.38,7.00,1.17,22,False,'D7E6EF',PP_ALIGN.CENTER)
    d.inventory.append({'page':len(d.pres.slides),'title':'请各位老师批评指正'})
    assert len(d.pres.slides)==24,len(d.pres.slides)
    output=OUT/'rebuilt_raw.pptx';d.pres.save(output)
    (OUT/'slide_plan.json').write_text(json.dumps(d.inventory,ensure_ascii=False,indent=2))
    sys.path.insert(0,str(ROOT/'中期/scripts'))
    from set_methods_fonts import set_fonts
    final=ROOT/'中期/谢东宸中期_逻辑重构版_v1.pptx'
    set_fonts(output,final)
    print(final)


if __name__=='__main__':main()
